"""Turn the video run sheet (markdown) into a Word document I can print or
read on a second screen while recording.

Usage:  python md_to_docx.py <in.md> <out.docx>
"""

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Arial"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark blue, matches the report headings
MUTED = RGBColor(0x60, 0x60, 0x60)


def set_run(run, size=11, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), FONT)


def add_runs(par, text, size=11, mono=False):
    """Handle **bold** and `code` inside a line."""
    text = text.replace("\\*", "*")
    text = re.sub(r"``([^`]+)``", r"`\1`", text)   # ``code`` -> `code`
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            inner = token[2:-2]
            if inner.startswith("`") and inner.endswith("`") and len(inner) > 2:
                inner = inner[1:-1]          # **bold** wrapped around `code`
            set_run(par.add_run(inner), size=size, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            set_run(par.add_run(token[1:-1]), size=size, bold=True,
                    color=ACCENT)
        else:
            set_run(par.add_run(token), size=size)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        pr.append(el)


def convert(src, dst):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # --- fenced code -------------------------------------------------
        if stripped.startswith("```"):
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.space_after = Pt(8)
            set_run(par.add_run("\n".join(block)), size=10, bold=False)
            continue

        # --- blank or a horizontal rule ----------------------------------
        if not stripped or set(stripped) <= {"-", "=", "*"}:
            i += 1
            continue

        # --- table -------------------------------------------------------
        if stripped.startswith("|") and i + 1 < n and re.match(
                r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t)
            for idx, text in enumerate(header):
                cell = t.rows[0].cells[idx]
                cell.text = ""
                par = cell.paragraphs[0]
                add_runs(par, text, size=10)
                for run in par.runs:
                    run.bold = True
                shade(cell, "E8EEF4")
            for row in rows:
                cells = t.add_row().cells
                for idx, text in enumerate(row):
                    cells[idx].text = ""
                    par = cells[idx].paragraphs[0]
                    add_runs(par, text.replace("<br>", " "), size=10)
            doc.add_paragraph()
            continue

        # --- headings ----------------------------------------------------
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            par.paragraph_format.space_after = Pt(6)
            size = {1: 17, 2: 14, 3: 12, 4: 11}[level]
            add_runs(par, m.group(2), size=size)
            for run in par.runs:
                run.bold = True
                run.font.color.rgb = ACCENT
            i += 1
            continue

        # --- blockquote ---------------------------------------------------
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(18)
            par.paragraph_format.space_after = Pt(6)
            add_runs(par, text, size=10)
            for run in par.runs:
                run.italic = True
                run.font.color.rgb = MUTED
            i += 1
            continue

        # --- bullets / numbered -------------------------------------------
        m = re.match(r"^([-*]|\d+\.)\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph(style="List Bullet" if m.group(1) in ("-", "*")
                                    else "List Number")
            add_runs(par, m.group(2), size=11)
            par.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # --- plain paragraph ----------------------------------------------
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(6)
        add_runs(par, stripped, size=11)
        i += 1

    doc.save(dst)
    return dst


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2]
    print("wrote", convert(src, dst))
