"""
make_storyboard.py - build the Giving Coast storyboard from the tutor's template.

The template (Storyboard-template.docx) is one screen: a "Storyboard" heading,
a TITLE / SCREEN ID / DATE line, an "Elements:" label over a sketch area, and
a five-row table (Description / Content / Interactions / Media Creation /
Navigation). This script fills that block once per screen of the real website
- seven screens, each with a live screenshot taken from the deployed site -
and saves the result as one document.

Formatting follows the template (Comic Sans MS, Letter, 2.54 cm margins), not
the A3 report, because this document is a design artifact, not the report.

Run:  python make_storyboard.py <template> <output>
"""

import copy
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

from storyboard_data import AUDIO_VIDEO, DATE, NOTES, SCREENS, TEXT_ITEMS

FONT = "Comic Sans MS"
BODY = 10

def set_cell_text(cell, text, size=BODY):
    """Write text into a table cell in the template font. Newlines become real
    line breaks inside one paragraph."""
    par = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for run in list(par.runs):
        run._element.getparent().remove(run._element)
    run = par.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)


def set_table_borders(table):
    """The template carries no 'Table Grid' style, so draw the borders
    directly on the table."""
    from docx.oxml import OxmlElement
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)


def boxed(doc, label, image=None, text=None, width_cm=16.5):
    """One bordered box with a Comic Sans label in its top-left corner -
    the visual equivalent of the template's GRAPHICS / TEXT / AUDIO-VIDEO
    frames. Text uses real line breaks between items."""
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    cell.width = Cm(width_cm)
    table.columns[0].width = Cm(width_cm)

    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run(label)
    r.font.name = FONT
    r.font.size = Pt(BODY)
    r.bold = False

    if image is not None:
        ip = cell.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.space_after = Pt(2)
        ip.add_run().add_picture(image, width=Cm(10.5))
        if text:
            cp = cell.add_paragraph()
            cp.paragraph_format.space_after = Pt(0)
            r = cp.add_run(text)
            r.font.name = FONT
            r.font.size = Pt(9)
    elif text:
        lines = text if isinstance(text, list) else [text]
        for line in lines:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            r.font.name = FONT
            r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_screen(doc, screen, img_dir, first=False):
    """Append one screen to the document, mirroring the template's layout."""
    if not first:
        br = doc.add_paragraph()
        br.paragraph_format.space_after = Pt(0)
        br.add_run().add_break(WD_BREAK.PAGE)

    def par(text="", size=BODY, bold=False, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.bold = bold
            rpr = r._element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                from docx.oxml import OxmlElement
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rfonts.set(qn(attr), FONT)
        return p

    # heading
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run("Storyboard")
    r.font.name = FONT
    r.font.size = Pt(14)
    r.bold = True

    par(f"TITLE: {screen['title']}      SCREEN ID: {screen['id']}      DATE: {DATE}")
    par("", space_after=0)

    par("Elements:", size=14, bold=True, space_after=2)

    # the template's three Elements frames: graphics / text / audio-video
    boxed(doc, "GRAPHICS:",
          image=f"{img_dir}/{screen['img']}",
          text="Live screenshot of the deployed page (https://24831941.it.scu.edu.au), taken with "
               "headless Chrome. The layout is the design; every value on it comes from the API.")
    boxed(doc, "TEXT:", text=TEXT_ITEMS[screen["id"]])
    boxed(doc, "AUDIO/VIDEO:", text=AUDIO_VIDEO[screen["id"]])

    par("NOTES", size=14, bold=True, space_after=2)
    boxed(doc, "NOTES:", text=NOTES[screen["id"]])

    table = doc.add_table(rows=5, cols=2)
    set_table_borders(table)
    labels = ["Description:", "Content:", "Interactions:", "Media Creation:", "Navigation:"]
    values = [screen["desc"], screen["content"], screen["interactions"],
              screen["media"], screen["navigation"]]
    for i, row in enumerate(table.rows):
        set_cell_text(row.cells[0], labels[i])
        set_cell_text(row.cells[1], values[i])
        row.cells[0].width = Cm(3.6)
        row.cells[1].width = Cm(12.9)
    for column, width in zip(table.columns, (Cm(3.6), Cm(12.9))):
        column.width = width


CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

URLS = {
    "S1-home.png": "/",
    "S2-search.png": "/search",
    "S3-detail.png": "/event/5",
    "S4-register.png": "/register?eventId=5",
    "S5-admin-list.png": "/admin/events",
    "S6-admin-new.png": "/admin/events/new",
    "S7-admin-edit.png": "/admin/events/5/edit",
}


CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

URLS = {
    "S1-home.png": "/",
    "S2-search.png": "/search",
    "S3-detail.png": "/event/5",
    "S4-register.png": "/register?eventId=5",
    "S5-admin-list.png": "/admin/events",
    "S6-admin-new.png": "/admin/events/new",
    "S7-admin-edit.png": "/admin/events/5/edit",
}


def take_screenshots(base_url, img_dir):
    """Re-shoot every screen from the deployed site.

    Headless Chrome rather than a screenshot library: no dependency to install,
    and --virtual-time-budget makes it wait for the API calls that fill the
    page. --timeout is not used because it captures the page before the data
    arrives, which is how the first attempt produced an empty admin list.
    """
    import subprocess
    os.makedirs(img_dir, exist_ok=True)
    for filename, path in URLS.items():
        target = os.path.join(img_dir, filename)
        subprocess.run(
            [CHROME, "--headless=new", "--disable-gpu", "--no-sandbox",
             "--hide-scrollbars", "--window-size=1280,1250",
             "--virtual-time-budget=30000",
             f"--screenshot={target}", base_url + path],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        print(f"  {filename} <- {base_url + path}")


def main(template_path, out_path, img_dir):
    doc = Document(template_path)
    # start from a clean body: the template's sample block is our layout
    # reference, but we rebuild every screen with the real content.
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)

    for i, screen in enumerate(SCREENS):
        add_screen(doc, screen, img_dir, first=(i == 0))
    doc.save(out_path)
    print(f"saved {out_path} with {len(SCREENS)} screens")


if __name__ == "__main__":
    # Usage:
    #   python tools/make_storyboard.py                 rebuild from existing screenshots
    #   python tools/make_storyboard.py --shoot         re-shoot the site, then rebuild
    #   python tools/make_storyboard.py --shoot --base https://24832054.it.scu.edu.au
    argv = sys.argv[1:]
    shoot = "--shoot" in argv
    base = "https://24831941.it.scu.edu.au"
    if "--base" in argv:
        base = argv[argv.index("--base") + 1]
        argv = argv[:argv.index("--base")] + argv[argv.index("--base") + 2:]
    positional = [a for a in argv if a != "--shoot"]

    template = positional[0] if len(positional) > 0 else "tools/Storyboard-template.docx"
    out = positional[1] if len(positional) > 1 else "submission/storyboard-giving-coast.docx"
    imgs = positional[2] if len(positional) > 2 else "submission/screenshots"

    if shoot:
        print(f"re-shooting 7 screens from {base}")
        take_screenshots(base, imgs)

    main(template, out, imgs)
